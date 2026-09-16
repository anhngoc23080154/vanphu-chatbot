"""Kiểm thử bộ đọc Excel bằng workbook nhỏ dựng trong bộ nhớ."""

from __future__ import annotations

import openpyxl
import pytest

from rag.ingest.sheets import parse_apartments, parse_qa, parse_shophouse
from rag.text import compress_ranges, norm_dir, parse_bedrooms, to_float, to_int


@pytest.fixture
def workbook():
    book = openpyxl.Workbook()
    book.remove(book.active)

    # --- Sheet căn hộ: header dòng 1, nhãn nhóm dòng 2, dữ liệu từ dòng 3 ---
    apartments = book.create_sheet("Bảng diện tích căn hộ")
    apartments.append(
        [
            "STT", "Mã Căn", "Mã Căn\n(Theo BVTĐ-GPXD)", "Mã Căn\nThương Mại", "Tầng",
            "Loại SP", "Tháp", "Tầng", None, "Căn", "Diện Tích\nTim Tường (m2)",
            "Diện Tích\nThông Thủy (m2)", "Số \nPhòng Ngủ", "Hướng \nBan Công",
            "Hướng Cửa", "Ghi Chú",
        ]
    )
    apartments.append(["2/ Căn Hộ Ở Tháp A"] + [None] * 15)
    apartments.append(
        ["1.0", "A-06-01", "A-NN-01", "CH-A06-01", "6.0", "CH", "A", "6.0", "1.0", "1.0",
         "102.33", "90.48", "3BR", "Nam", "Bắc", None]
    )
    apartments.append(
        ["2.0", "A-06-02", "A-NN-02", "CH-A06-02", "6.0", "CH", "A", "6.0", "2.0", "2.0",
         "77.61", "69.73", "2BR", "Tây", "ĐÔng", None]
    )
    # Penthouse: dòng chính + dòng phụ tầng 35 + dòng tổng 34+35.
    apartments.append(
        ["3.0", "A-34-01", "A-NN-03", "PH-A34-03", "34.0", "PH", "A", "34.0", "1.0", "3.0",
         "166.0", "151.83", "3BR+1", "Tây Nam", "Đông Nam", None]
    )
    apartments.append([None] * 4 + ["35.0", None, None, "35.0", None, None, "141.06", "123.67"] + [None] * 4)
    apartments.append([None] * 4 + ["34+35", None, None, "34+35", None, None, "307.06", "275.50"] + [None] * 4)

    # --- Sheet shophouse: tiêu đề dòng 1-2, header dòng 4, dữ liệu từ dòng 6 ---
    shop = book.create_sheet("diện tích shophouse")
    shop.append(["QUY CÁCH THỐNG KÊ MÃ SẢN PHẨM (D6-D8-L9)"] + [None] * 13)
    shop.append(["Dự án: 03HNCM008"] + [None] * 13)
    shop.append([None] * 14)
    shop.append(
        ["STT", "Mã Căn\n(Theo BVTĐ-GPXD)", "Mã Căn\nThương Mại", "Tầng", "Loại SP", "Tháp",
         "Tầng", "Căn", "Diện Tích\nTim Tường (m2)", "Diện Tích\nThông Thủy (m2)",
         "Số \nPhòng Ngủ", "Hướng \nBan Công", "Hướng Cửa", "Ghi Chú"]
    )
    shop.append(["1/ Shophouse & TMDV"] + [None] * 13)
    shop.append(["1.0", "A-01-01", "SH-A01-01", "1.0", "SH", "A", "1.0", "1.0", "83.88", "77.61"] + [None] * 4)
    shop.append([None, None, None, "2.0", None, None, "2.0", None, "95.45", "87.98"] + [None] * 4)
    shop.append([None, None, None, "1+2", None, None, "1+2", None, "179.33", "165.59"] + [None] * 4)
    shop.append(["2.0", "A-03-01", "TM-A03-09", "3.0", "TM", "A", "3.0", "9.0", "79.6", "73.8"] + [None] * 4)

    # --- Sheet Q&A: tiêu đề dòng 2, header dòng 3 ---
    qa = book.create_sheet("Q&A ")
    qa.append([None] * 5)
    qa.append(["BỘ NGÂN HÀNG Q&A VLASTA PREMIER PHÚ THUẬN", None, None, None, None])
    qa.append(["STT", "CÂU HỎI ", "TRẢ LỜI", None, None])
    qa.append([None, "THÔNG TIN LIÊN QUAN ĐẾN CHỦ ĐẦU TƯ VÀ ĐỐI TÁC", None, None, None])
    qa.append(["1.0", "Tên dự án", "Vlasta Premier – Phú Thuận", None, None])
    qa.append(["2", "Chủ đầu tư", "Công ty Cổ phần Đầu tư Xây dựng New Tech", None, "ghi chú"])
    qa.append(["3", "Câu hỏi chưa có đáp án", None, None, None])
    qa.append([None, "KẾT NỐI VÙNG", None, None, None])
    qa.append(["4", "Trường học lân cận?", "Đại học RMIT Nam Sài Gòn", None, None])
    qa.append(
        ["Dù các tài liệu mô tả về dự án với sự cẩn trọng cao, chúng tôi không đảm bảo tính "
         "chính xác hoàn toàn và miễn trách nhiệm pháp lý đối với nội dung các tài liệu này. "
         "Thông tin có thể thay đổi theo chấp thuận của cơ quan Nhà Nước có thẩm quyền.",
         None, None, None, None]
    )
    return book


def test_doc_can_ho_bo_qua_dong_nhan_nhom(workbook):
    units = parse_apartments(workbook)
    primary = [unit for unit in units if unit.is_primary]
    assert len(primary) == 3
    assert [unit.product_type for unit in primary] == ["CH", "CH", "PH"]


def test_can_ho_dau_tien_dung_gia_tri(workbook):
    unit = parse_apartments(workbook)[0]
    assert unit.unit_code == "A-06-01"
    assert unit.unit_code_alt == "A-NN-01"
    assert unit.unit_code_commercial == "CH-A06-01"
    assert unit.tower == "A"
    assert unit.floor == 6
    assert unit.floor_label == "6"
    assert unit.unit_no == "01"
    assert unit.area_gross == pytest.approx(102.33)
    assert unit.area_net == pytest.approx(90.48)
    assert unit.bedroom_count == 3
    assert unit.balcony_dir == "Nam"
    assert unit.door_dir == "Bắc"


def test_sua_loi_go_dau_huong(workbook):
    unit = parse_apartments(workbook)[1]
    assert unit.door_dir == "Đông"


def test_penthouse_ke_thua_ma_tu_dong_chinh(workbook):
    units = parse_apartments(workbook)
    penthouse = [unit for unit in units if unit.product_type == "PH"]
    assert len(penthouse) == 3
    assert all(unit.unit_code == "A-34-01" for unit in penthouse)
    assert all(unit.unit_code_commercial == "PH-A34-03" for unit in penthouse)
    assert [unit.floor_label for unit in penthouse] == ["34", "35", "34+35"]
    assert [unit.is_primary for unit in penthouse] == [True, False, False]
    assert [unit.is_total_row for unit in penthouse] == [False, False, True]
    assert penthouse[2].area_gross == pytest.approx(307.06)


def test_doc_shophouse_lech_cot(workbook):
    units = parse_shophouse(workbook)
    shophouse = [unit for unit in units if unit.product_type == "SH"]
    assert len(shophouse) == 3
    assert shophouse[0].unit_code == "A-01-01"
    assert shophouse[0].unit_code_commercial == "SH-A01-01"
    assert shophouse[0].area_gross == pytest.approx(83.88)
    assert shophouse[2].is_total_row and shophouse[2].floor_label == "1+2"

    tmdv = [unit for unit in units if unit.product_type == "TM"]
    assert len(tmdv) == 1 and tmdv[0].unit_code_commercial == "TM-A03-09"


def test_doc_qa_nhan_dung_nhom_va_bo_cau_thieu_dap_an(workbook):
    items = parse_qa(workbook)
    stts = [item.stt for item in items]
    assert stts == ["1", "2", "4", "disclaimer"]
    assert items[0].section == "THÔNG TIN LIÊN QUAN ĐẾN CHỦ ĐẦU TƯ VÀ ĐỐI TÁC"
    assert items[2].section == "KẾT NỐI VÙNG"
    assert "ghi chú" in items[1].note
    assert items[3].section == "Miễn trừ trách nhiệm"


# --- Hàm tiện ích xử lý chuỗi/số ---
@pytest.mark.parametrize(
    "value,expected",
    [("102.33", 102.33), ("102,33", 102.33), ("1.234,56", 1234.56), ("", None), ("1+2", None)],
)
def test_to_float(value, expected):
    result = to_float(value)
    if expected is None:
        assert result is None
    else:
        assert result == pytest.approx(expected)


def test_to_int_bo_qua_gia_tri_gop_tang():
    assert to_int("6.0") == 6
    assert to_int("34+35") is None


def test_norm_dir_sua_loi_go():
    assert norm_dir("ĐÔng ") == "Đông"
    assert norm_dir("tây nam") == "Tây Nam"


def test_parse_bedrooms():
    assert parse_bedrooms("3BR") == ("3BR", 3)
    assert parse_bedrooms("1BR+1") == ("1BR+1", 1)
    assert parse_bedrooms(None) == ("", None)


def test_compress_ranges():
    assert compress_ranges([6, 7, 8, 10, 11, 20]) == "6-8, 10-11, 20"
    assert compress_ranges([]) == ""
